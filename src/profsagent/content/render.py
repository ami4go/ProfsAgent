"""Render a lecture deck as .pptx (python-pptx, 16:9) with speaker notes, and as a Markdown handout.

Layout: a title slide, an objectives slide, the content slides (bullets on the left, the figure on the right; a slide
with only a figure uses the full width) and a references slide. Figures are native shapes, so the professor can edit
them in PowerPoint: a flow is a column of boxes joined by arrows, code a dark panel in a monospace font, a table a
PowerPoint table. A check slide shows its question with lettered options; the answer and its explanation go to the
speaker notes only.

python-pptx cannot measure text, so every text box is fitted by estimate (`fit`: average glyph width and line height
for the font size): the size steps down until the text fits, and text that does not fit even at the smallest size is
reported as an overflow (a D-OVERFLOW warning in material.json).
"""
from __future__ import annotations

import math
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from profsagent.content.models import Figure, LectureDeck, Quiz, Slide

BADGE = {"example": "Example", "worked_problem": "Worked problem", "activity": "Activity", "check": "Check", "summary": "Summary"}
def answer(q: Quiz) -> str:
    """The answer line for notes and the handout ("?" when the index is invalid; D-QUIZ already reports that)."""
    ok = 0 <= q.answer < len(q.options)
    return f"Answer: {'ABCDE'[q.answer]}. {q.options[q.answer]}" if ok else "Answer: ? (invalid answer index)"


GLYPH, LEADING, MONO_GLYPH = 0.52, 1.2, 0.55  # average glyph width (Calibri; Consolas is exactly 0.55) and line height,
                                              # as fractions of the font size
TEXT_SHARE = {"code": 0.42}                   # bullets' share of the width beside a figure (code needs the room); else 0.52


def fit(paragraphs: list[str], w, h, start: int, smallest: int, gap_pt: float = 0.0, glyph: float = GLYPH) -> tuple[int, bool]:
    """The largest font size from `start` down to `smallest` at which the paragraphs fit a w x h (EMU) box, and whether
    they fit at all. `gap_pt`: space after each paragraph."""
    width_pt, height_pt = Emu(w).pt - 8, Emu(h).pt - 8          # text-frame margins
    for size in range(start, smallest - 1, -1):
        per_line = max(1, int(width_pt / (glyph * size)))
        lines = sum(max(1, math.ceil(len(p) / per_line)) for p in paragraphs)
        if lines * size * LEADING + gap_pt * max(0, len(paragraphs) - 1) <= height_pt:
            return size, True
    return smallest, False


class Deck:
    def __init__(self, cfg: dict, footer: str):
        r = cfg["render"]
        self.c = {k: RGBColor.from_string(x) for k, x in r["colors"].items()}
        self.font, self.mono, self.footer = r["font"], r["mono_font"], footer
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(r["width_in"]), Inches(r["height_in"])
        self.W, self.H = self.prs.slide_width, self.prs.slide_height
        self.blank = self.prs.slide_layouts[6]
        self.n = 0
        self.overflows: list[dict] = []          # {slide, what}: text that does not fit even at the smallest size

    def _fitted(self, what: str, paragraphs: list[str], w, h, start: int, smallest: int, gap_pt: float = 0.0,
                glyph: float = GLYPH) -> int:
        size, ok = fit(paragraphs, w, h, start, smallest, gap_pt, glyph)
        if not ok:
            self.overflows.append({"slide": self.n, "what": what})
        return size

    # ------------------------------------------------------------------ primitives
    def _text(self, slide, x, y, w, h, text: str, size: int, color: str = "ink", bold=False, font=None,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap, tf.vertical_anchor = True, anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            run = p.add_run()
            run.text = line
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, font or self.font
            run.font.color.rgb = self.c[color]
        return tb

    def _box(self, slide, shape, x, y, w, h, fill: str, line: str | None = None):
        s = slide.shapes.add_shape(shape, x, y, w, h)
        s.fill.solid()
        s.fill.fore_color.rgb = self.c[fill]
        if line:
            s.line.color.rgb = self.c[line]
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        return s

    def _new(self, notes: str = ""):
        slide = self.prs.slides.add_slide(self.blank)
        self.n += 1
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def _chrome(self, slide, title: str, badge: str | None = None, sources: list[str] | None = None):
        m = Inches(0.6)
        self._box(slide, MSO_SHAPE.RECTANGLE, 0, Inches(0.45), Inches(0.14), Inches(0.75), "accent")
        tw, th = self.W - 2 * m - (Inches(2.2) if badge else 0), Inches(0.95)
        self._text(slide, m, Inches(0.35), tw, th, title, self._fitted("title", [title], tw, th, 30, 20),
                   bold=True, anchor=MSO_ANCHOR.MIDDLE)
        if badge:
            b = self._box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, self.W - m - Inches(2.0), Inches(0.55), Inches(2.0), Inches(0.5),
                          "accent_soft")
            tf = b.text_frame
            tf.text = badge
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.runs[0].font.size, p.runs[0].font.bold, p.runs[0].font.name = Pt(14), True, self.font
            p.runs[0].font.color.rgb = self.c["accent"]
        foot = self.footer + (f"   ·   Sources: {' '.join(f'[{s}]' for s in sources)}" if sources else "")
        self._text(slide, m, self.H - Inches(0.5), self.W - 2 * m - Inches(0.8), Inches(0.35), foot, 11, "muted")
        self._text(slide, self.W - m - Inches(0.8), self.H - Inches(0.5), Inches(0.8), Inches(0.35), str(self.n), 11, "muted",
                   align=PP_ALIGN.RIGHT)

    def _bullets(self, slide, x, y, w, h, bullets: list[str], start: int = 24):
        """Real PowerPoint bullets (hanging indent, so wrapped lines align with the text)."""
        indent = Inches(0.32)
        size = self._fitted("bullets", bullets, w - indent, h, start, 14, gap_pt=10)
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(indent))
            pPr.set("indent", str(-indent))
            clr = etree.SubElement(etree.SubElement(pPr, qn("a:buClr")), qn("a:srgbClr"))
            clr.set("val", str(self.c["accent"]))
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "▪")
            run = p.add_run()
            run.text = b
            run.font.size, run.font.name = Pt(size), self.font
            run.font.color.rgb = self.c["ink"]

    # ------------------------------------------------------------------ figures
    def _figure(self, slide, f: Figure, x, y, w, h):
        cap_h = Inches(0.4) if f.caption else 0
        if f.kind == "flow":
            self._flow(slide, f.steps, x, y, w, h - cap_h)
        elif f.kind == "code":
            self._code(slide, f, x, y, w, h - cap_h)
        elif f.kind == "table":
            self._table(slide, f.rows, x, y, w, h - cap_h)
        if f.caption:
            self._text(slide, x, y + h - cap_h, w, cap_h, f.caption, 12, "muted", align=PP_ALIGN.CENTER)

    def _flow(self, slide, steps: list[str], x, y, w, h):
        steps = [s for s in steps if s.strip()]
        gap = Inches(0.35)
        bh = Emu(int(min(Inches(0.95), (h - gap * (len(steps) - 1)) / len(steps))))
        top = y + (h - (bh * len(steps) + gap * (len(steps) - 1))) // 2
        size = min(self._fitted("flow step", [s], w, bh, 16, 11) for s in steps)
        for i, s in enumerate(steps):
            by = top + i * (bh + gap)
            b = self._box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, by, w, bh, "accent_soft", "accent")
            tf = b.text_frame
            tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
            tf.text = s
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.runs[0].font.size, p.runs[0].font.name = Pt(size), self.font
            p.runs[0].font.color.rgb = self.c["ink"]
            if i < len(steps) - 1:
                self._box(slide, MSO_SHAPE.DOWN_ARROW, x + w // 2 - Inches(0.15), by + bh + Inches(0.03), Inches(0.3),
                          gap - Inches(0.06), "accent")

    def _code(self, slide, f: Figure, x, y, w, h):
        """Code never wraps (a wrapped line reads as two statements): the size is the largest at which the longest line
        and every line fit; below 9 pt the panel overflows and is reported."""
        panel = self._box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, "code_bg")
        panel.adjustments[0] = 0.04
        lines = f.code.rstrip("\n").expandtabs(4).split("\n")
        label = f"{f.language}\n" if f.language else ""
        iw, ih = w - Inches(0.4), h - Inches(0.3)
        rows, longest = len(lines) + (1 if f.language else 0), max(len(ln) for ln in lines)
        size = next((s for s in range(14, 8, -1) if longest * MONO_GLYPH * s <= Emu(iw).pt - 8 and rows * s * LEADING <= Emu(ih).pt - 8), 9)
        if size == 9 and (longest * MONO_GLYPH * 9 > Emu(iw).pt - 8 or rows * 9 * LEADING > Emu(ih).pt - 8):
            self.overflows.append({"slide": self.n, "what": "code"})
        tb = self._text(slide, x + Inches(0.2), y + Inches(0.15), iw, ih, label + "\n".join(lines), size, "code_ink", font=self.mono)
        tb.text_frame.word_wrap = False
        if f.language:
            r = tb.text_frame.paragraphs[0].runs[0]
            r.font.color.rgb, r.font.size = self.c["accent_soft"], Pt(11)

    def _table(self, slide, rows: list[list[str]], x, y, w, h):
        t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, min(h, Inches(0.55) * len(rows))).table
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.text = val
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.c["accent"] if i == 0 else self.c["panel"]
                for p in cell.text_frame.paragraphs:
                    for r_ in p.runs:
                        r_.font.size, r_.font.name, r_.font.bold = Pt(14), self.font, i == 0
                        r_.font.color.rgb = self.c["accent_soft"] if i == 0 else self.c["ink"]

    # ------------------------------------------------------------------ slides
    def title_slide(self, course: str, title: str, subtitle: str, note: str):
        s = self._new()
        self._box(s, MSO_SHAPE.RECTANGLE, 0, 0, self.W, self.H, "accent")
        self._text(s, Inches(0.9), Inches(1.4), self.W - Inches(1.8), Inches(0.5), course, 18, "accent_soft")
        tw, th = self.W - Inches(1.8), Inches(2.2)
        self._text(s, Inches(0.9), Inches(2.0), tw, th, title, self._fitted("title", [title], tw, th, 44, 28), "accent_soft",
                   bold=True, anchor=MSO_ANCHOR.MIDDLE)
        self._text(s, Inches(0.9), Inches(4.4), self.W - Inches(1.8), Inches(0.6), subtitle, 18, "accent_soft")
        self._text(s, Inches(0.9), self.H - Inches(0.9), self.W - Inches(1.8), Inches(0.4), note, 12, "accent_soft")

    def list_slide(self, title: str, items: list[str], notes: str = "", start: int = 24):
        s = self._new(notes)
        self._chrome(s, title)
        self._bullets(s, Inches(0.6), Inches(1.6), self.W - Inches(1.2), self.H - Inches(2.3), items, start)

    def _quiz(self, slide, q: Quiz, x, y, w, h):
        qh = Inches(1.3)
        self._text(slide, x, y, w, qh, q.question, self._fitted("question", [q.question], w, qh, 26, 16), bold=True)
        top, gap = y + qh + Inches(0.2), Inches(0.18)
        oh = Emu(int(min(Inches(0.85), (h - qh - Inches(0.2) - gap * (len(q.options) - 1)) / len(q.options))))
        size = min(self._fitted("option", [o], w - Inches(1.1), oh, 20, 12) for o in q.options)
        for i, opt in enumerate(q.options):
            oy = top + i * (oh + gap)
            self._box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, oy, w, oh, "panel")
            letter = self._box(slide, MSO_SHAPE.OVAL, x + Inches(0.15), oy + (oh - Inches(0.5)) // 2, Inches(0.5), Inches(0.5), "accent")
            letter.text_frame.text = "ABCDE"[i]
            r = letter.text_frame.paragraphs[0]
            r.alignment = PP_ALIGN.CENTER
            r.runs[0].font.size, r.runs[0].font.bold, r.runs[0].font.name = Pt(16), True, self.font
            r.runs[0].font.color.rgb = self.c["accent_soft"]
            self._text(slide, x + Inches(0.85), oy, w - Inches(1.1), oh, opt, size, anchor=MSO_ANCHOR.MIDDLE)

    def content_slide(self, sl: Slide):
        notes = sl.notes
        if sl.quiz:
            q = sl.quiz
            notes += f"\n\n{answer(q)}\n{q.explanation}".rstrip()
        s = self._new(notes)
        self._chrome(s, sl.title, BADGE.get(sl.kind), sl.sources)
        m, top = Inches(0.6), Inches(1.6)
        body_h = self.H - top - Inches(0.75)
        has_fig = sl.figure.kind != "none"
        if sl.quiz:
            self._quiz(s, sl.quiz, m + Inches(0.8), top, self.W - 2 * m - Inches(1.6), body_h)
        elif has_fig and sl.bullets:
            left_w = int((self.W - 2 * m) * TEXT_SHARE.get(sl.figure.kind, 0.52))
            self._bullets(s, m, top, left_w, body_h, sl.bullets)
            self._figure(s, sl.figure, m + left_w + Inches(0.4), top, self.W - 2 * m - left_w - Inches(0.4), body_h)
        elif has_fig:
            self._figure(s, sl.figure, m + Inches(1.2), top, self.W - 2 * m - Inches(2.4), body_h)
        else:
            self._bullets(s, m, top, self.W - 2 * m, body_h, sl.bullets)

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)


def render_pptx(deck: LectureDeck, frame: dict, sources: dict[str, dict], cfg: dict, path: Path) -> tuple[int, list[dict]]:
    """`frame`: title-slide and objectives text decided by code (generate.frame). Returns the number of slides and the
    overflows ({slide, what}: text that does not fit even at the smallest size)."""
    d = Deck(cfg, frame["footer"])
    d.title_slide(frame["course"], deck.title, frame["subtitle"], frame["draft_note"])
    d.list_slide("By the end of this lecture you can", frame["objectives"], frame["objectives_notes"])
    for sl in deck.slides:
        d.content_slide(sl)
    cited = [sid for sid in sources if any(sid in sl.sources for sl in deck.slides)]
    if cited:
        d.list_slide("References", [f"[{sid}] {ref_line(sources[sid])}" for sid in cited], start=18)
    d.save(path)
    return d.n, d.overflows


def ref_line(src: dict) -> str:
    year = f" ({src['year']})" if src.get("year") else ""
    return f"{src['title']}{year}" + (f". {src['url']}" if src.get("url") else "")


def render_markdown(deck: LectureDeck, frame: dict, sources: dict[str, dict]) -> str:
    """The handout: every slide with its notes, then the references."""
    L = [f"# {deck.title}", "", f"*{frame['course']} · {frame['subtitle']}*", "", f"> {frame['draft_note']}", "",
         "## Learning outcomes", "", *[f"- {o}" for o in frame["objectives"]], ""]
    for i, sl in enumerate(deck.slides, 1):
        L += [f"## {i}. {sl.title}" + (f" ({BADGE[sl.kind]})" if sl.kind in BADGE else ""), ""]
        L += [f"- {b}" for b in sl.bullets]
        f = sl.figure
        if f.kind == "flow":
            L += ["", " → ".join(f.steps)]
        elif f.kind == "code":
            L += ["", f"```{f.language}", f.code.rstrip(), "```"]
        elif f.kind == "table":
            L += ["", "| " + " | ".join(f.rows[0]) + " |", "|" + "---|" * len(f.rows[0])] + ["| " + " | ".join(r) + " |" for r in f.rows[1:]]
        if f.caption:
            L += ["", f"*{f.caption}*"]
        if sl.quiz:
            q = sl.quiz
            L += ["", f"**{q.question}**", ""] + [f"{'ABCDE'[i]}. {o}" for i, o in enumerate(q.options)]
            L += ["", f"*{answer(q)}. {q.explanation}*"]
        meta = [f"{sl.minutes:g} min", "outcomes " + ", ".join(sl.los) if sl.los else "", "sources " + " ".join(f"[{s}]" for s in sl.sources) if sl.sources else ""]
        L += ["", f"**Notes** ({' · '.join(x for x in meta if x)}): {sl.notes}", ""]
    cited = [sid for sid in sources if any(sid in sl.sources for sl in deck.slides)]
    if cited:
        L += ["## References", ""] + [f"- [{sid}] {ref_line(sources[sid])}" for sid in cited]
    return "\n".join(L) + "\n"
