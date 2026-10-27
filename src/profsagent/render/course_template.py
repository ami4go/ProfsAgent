"""The official IIIT-D CSE course description template (templates/iiitd_cse_course_template.xlsx), filled from a run.

`template_data(state)` maps the checked design onto the template's fields (no LLM). `write_xlsx` fills a copy of the
official workbook: every label row, both "Do not remove or edit this row" markers, the Taxonomy sheet, the Department
drop-down and the programme outcome texts are kept exactly; only the dummy/placeholder rows are replaced by real rows,
and one mapping + explanation column pair is added per CO. The PDF export (scripts/export_course_sheet_pdf.py) prints
the same `template_data`, so both documents always agree.
"""
from __future__ import annotations

import re
from collections import defaultdict
from copy import copy
from pathlib import Path

from profsagent.config import ROOT

TEMPLATE = ROOT / "templates" / "iiitd_cse_course_template.xlsx"
TYPE_LABEL = {"quiz": "Quiz", "assignment": "Assignment", "midsem": "Mid-sem exam", "endsem": "End-sem exam", "project": "Project",
              "report": "Report", "lab": "Lab", "presentation": "Presentation", "participation": "Participation"}
# the template's "Assignment/lab/tutorial" column: take-home and practical work only; quizzes and exams are in the assessment plan
ACTIVITY_TYPES = {"assignment", "lab", "project", "report", "presentation"}
ROLE_TYPE = {"primary": "Textbook", "reference": "Reference", "reading": "Reference (paper)"}
PROGRAMME_LABEL = {"MTECH-CSE": "MTech(CSE)", "BTECH-COMMON": "B.Tech. (common)", "BTECH-CSE": "B.Tech.(CSE)", "BTECH-CSAI": "B.Tech.(CSAI)"}


def pretty_assessment(iid: str) -> str:
    s = iid.split(":", 1)[-1]
    s = re.sub(r"^ProjMilestone", "Project Milestone ", s)
    s = re.sub(r"^ProjFinal", "Project (final)", s)
    s = re.sub(r"^Proj(?!ect)", "Project ", s)
    s = s.replace("Midsem", "Mid-sem exam").replace("Endsem", "End-sem exam")
    s = re.sub(r"(?<=[a-z])(?=\d)", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _is_proctored(c: dict, types: list[str]) -> bool:
    return c["proctored"] if c.get("proctored") is not None else str(c["type"]).lower() in types


def template_data(st: dict) -> dict:
    from profsagent.config import pedagogy
    ped = pedagogy()
    cc, f, d = st["cco_compact"], st["cco"]["fields"], st["defaults"]
    lo_parent = {lo["id"]: lo["parent_co"] for lo in st.get("los", [])}
    pos = st.get("positioning", {})
    pre = defaultdict(list)
    for p in pos.get("prerequisites", []):
        alts = [a for a in p.get("alternatives", []) if a != p["course_uid"]]
        pre[p["kind"]].append(p["course_uid"] + (f" (or {', '.join(alts)})" if alts else ""))
    code_re = re.compile(r"[A-Z]{2,4}\s?\d{3}")
    other = [b for b in (cc.get("target_students") or {}).get("assumed_background") or [] if not code_re.search(b)]
    level = st.get("level") or d.get("level") or 3
    progs = set((f.get("target_students") or {}).get("programmes") or [])
    offered = "PG" if int(level) >= 5 else ("UG, PG" if any(p.startswith("MTECH") for p in progs) else "UG")
    header = [("Code", (f.get("level_or_code") or {}).get("code") or ""), ("Department", st.get("department", "CSE")),
              ("Name", cc["course_name"]), ("Credits", d["credits"]), ("OfferedTo", offered),
              ("Description", (st.get("narrative") or {}).get("course_description", "")),
              ("Mandatory Prereg", ", ".join(pre["mandatory"]) or "None"), ("Desired Prereq", ", ".join(pre["desirable"]) or "None"),
              ("Other Prereq", "; ".join(other) or "None"),
              ("Antireq", ", ".join(a["course_uid"] for a in pos.get("anti_requisites", [])) or "None")]
    cos = [{"id": c["id"], "statement": c["statement"], "bl": f"BL{int(c['bloom_level'])}"} for c in st.get("cos", [])]

    # activities per topic: assessments, labs and tutorials that use it
    acts = defaultdict(list)
    for comp in (st.get("assessment") or {}).get("components", []):
        if str(comp.get("type")).lower() not in ACTIVITY_TYPES:
            continue
        for i in comp["instances"]:
            for t in i.get("covers_topics", []):
                acts[t].append(pretty_assessment(i["id"]))
    lt = st.get("labs_tutorials") or {}
    for s in lt.get("labs", []):
        for t in s.get("uses_topics", []):
            acts[t].append(f"Lab (week {s['week']})")
    for s in lt.get("tutorials", []):
        for t in s.get("uses_topics", []):
            acts[t].append(f"Tutorial (week {s['week']})")
    topics = []
    for m in st.get("modules", []):
        for t in m["topics"]:
            co_ids = sorted({lo_parent.get(lo, "") for lo in t.get("serves_los", [])} - {""})
            topics.append({"title": t["title"], "cos": ", ".join(co_ids), "activities": "; ".join(dict.fromkeys(acts.get(t["id"], []))),
                           "load": t.get("outside_load_hours")})
    labs = []
    for s in sorted(lt.get("labs", []), key=lambda x: x["week"]):
        co_ids = sorted({lo_parent.get(x, "") for x in s.get("practises_los", [])} - {""})
        labs.append({"title": f"Week {s['week']}: " + (s.get("title") + ": " if s.get("title") else "") + s["exercise"],
                     "cos": ", ".join(co_ids), "platform": ", ".join(s.get("tools", []))})
    ptypes = ped.get("proctored", {}).get("types", [])
    assess = []
    for c in (st.get("assessment") or {}).get("components", []):
        n = len(c["instances"])
        label = TYPE_LABEL.get(str(c["type"]).lower(), str(c["type"]).title())
        if n > 1:
            label += f" ({n}" + (f", best {c['best_k']}" if c.get("best_k") else "") + ")"
        if _is_proctored(c, ptypes):
            label += " [proctored]"
        assess.append({"type": label, "weight": c["weight_pct"], "proctored": _is_proctored(c, ptypes)})
    resources = []
    for r in st.get("resources", []):
        ident = r["identifier"].get("isbn") and f"ISBN {r['identifier']['isbn']}" or r["identifier"].get("doi") and f"DOI {r['identifier']['doi']}" or ""
        rtype = ROLE_TYPE.get(r["role"], r["role"].title())
        if not r["identifier"].get("isbn") and not r["identifier"].get("doi") and r["identifier"].get("url"):
            rtype = "Online"
        resources.append({"title": f"{', '.join(r['authors'])} ({r.get('year')}). {r['title']}." + (f" {ident}" if ident else ""), "type": rtype})

    # CO -> programme outcome mapping: computed strength when the assessment plan evidences it, else the provisional one
    computed = {(r["co_id"], r["po_uid"]): r["strength_computed"] for r in st.get("copo_computed", [])}
    mapping = {}
    for m in st.get("copo", []):
        s = computed.get((m["co_id"], m["po_uid"])) or m["strength_provisional"]
        mapping[(m["co_id"], m["po_uid"])] = (int(s), m.get("justification", ""))
    return {"header": header, "cos": cos, "approval_date": "", "proposing_dept": st.get("department", "CSE"),
            "topics": topics, "labs": labs, "assessment": assess, "resources": resources, "mapping": mapping,
            "proctored_pct": sum(a["weight"] for a in assess if a["proctored"])}


# ------------------------------------------------------------------ xlsx
def write_xlsx(st: dict, out: Path) -> Path:
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    data = template_data(st)
    wb = openpyxl.load_workbook(TEMPLATE)
    src = wb["Template"]
    idx = wb.sheetnames.index("Template")
    maxc = 12
    proto = {r: [src.cell(r, c) for c in range(1, maxc + 1)] for r in range(1, src.max_row + 1)}
    heights = {r: src.row_dimensions[r].height for r in proto}
    widths = {get_column_letter(c): src.column_dimensions[get_column_letter(c)].width for c in range(1, maxc + 1)}
    dv_formula = next((d.formula1 for d in src.data_validations.dataValidation), '"CSE,ECE,SSH,Math,HCD,CB"')
    ws = wb.create_sheet("Template_new", idx)
    for col, wdt in widths.items():
        if wdt:
            ws.column_dimensions[col].width = wdt
    row = 0
    merges = []

    def emit(proto_row: int, values: dict[int, object] | None = None, keep: set[int] | None = None):
        """Write one row styled like `proto_row`; copy the prototype's text in columns `keep`, then apply `values`."""
        nonlocal row
        row += 1
        for c, pc in enumerate(proto[proto_row], start=1):
            cell = ws.cell(row, c)
            if pc.has_style:
                cell.font, cell.border, cell.fill = copy(pc.font), copy(pc.border), copy(pc.fill)
                cell.number_format, cell.alignment, cell.protection = pc.number_format, copy(pc.alignment), copy(pc.protection)
            if keep and c in keep:
                cell.value = pc.value
        for c, val in (values or {}).items():
            ws.cell(row, c).value = val
        if heights.get(proto_row):
            ws.row_dimensions[row].height = heights[proto_row]
        return row

    ALL = set(range(1, maxc + 1))
    for i, (_, val) in enumerate(data["header"], start=1):             # rows 1-10
        r = emit(i, {2: val}, keep={1})
        if i == 6:
            merges.append(f"B{r}:D{r}")
    emit(11, keep={1, 2})                                              # CourseOutcomes (+ guidance)
    emit(12, keep=ALL - {11, 12})                                      # CO # | Outcome | BL # | BL guidance | Taxonomy
    for c in data["cos"]:
        emit(13, {1: c["id"], 2: c["statement"], 3: c["bl"]})
    emit(16, {2: data["approval_date"]}, keep={1})                     # Approval Date (empty for a new course)
    dept_row = emit(17, {2: data["proposing_dept"]}, keep={1})
    r = emit(18, keep={1, 2}); merges.append(f"B{r}:C{r}")             # EndMandatory: verbatim
    emit(19, keep={1}); emit(20, keep=ALL)
    for t in data["topics"]:
        emit(21, {1: t["title"], 2: t["cos"], 3: t["activities"], 4: t["load"]})
    emit(22, keep={1}); emit(23, keep=ALL)
    for lab in data["labs"]:
        emit(24, {1: lab["title"], 2: lab["cos"], 3: lab["platform"]})
    emit(25, keep={1}); emit(26, keep=ALL)
    for a in data["assessment"]:
        emit(27, {1: a["type"], 2: a["weight"]})
    emit(28, keep={1}); emit(29, keep=ALL)
    for res in data["resources"]:
        emit(30, {1: res["title"], 2: res["type"]})
    r = emit(31, keep={1, 2}); merges.append(f"B{r}:C{r}")             # EndSuggested: verbatim
    r = emit(32, keep={1}); merges.append(f"A{r}:I{r}")
    # CO -> PO mapping header: one (COk, COk-Explanation) column pair per CO
    co_cols = {c["id"]: 4 + 2 * k for k, c in enumerate(data["cos"])}
    r = emit(33, keep={1, 2}); merges.append(f"B{r}:C{r}")
    for co, col in co_cols.items():
        for cc_, txt in ((col, co), (col + 1, f"{co}-Explanation")):
            cell = ws.cell(r, cc_, txt)
            cell.font, cell.alignment, cell.border = copy(proto[33][3].font), copy(proto[33][3].alignment), copy(proto[33][3].border)
    r = emit(34, keep={1, 2}); merges.append(f"B{r}:C{r}")
    for c in data["cos"]:
        ws.cell(r, co_cols[c["id"]], c["statement"]).alignment = copy(proto[34][1].alignment)
    block_start, cur = None, None
    for pr in range(35, 71):
        label = proto[pr][0].value
        if label:
            if cur:
                merges.append(f"A{block_start}:A{row}")
            cur = {v: k for k, v in PROGRAMME_LABEL.items()}.get(label.strip())
            block_start = row + 1
        r = emit(pr, keep={1, 2, 3})
        code = str(proto[pr][1].value).strip()
        for co, col in co_cols.items():
            hit = data["mapping"].get((co, f"{cur}/{code}"))
            if hit:
                ws.cell(r, col, hit[0])
                ws.cell(r, col + 1, hit[1]).alignment = copy(proto[pr][2].alignment)
    merges.append(f"A{block_start}:A{row}")
    for col in range(4, 4 + 2 * len(co_cols)):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 8 if col % 2 == 0 else 40
    for m in merges:
        ws.merge_cells(m)
    dv = DataValidation(type="list", formula1=dv_formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B2")
    dv.add(f"B{dept_row}")
    wb.remove(src)
    ws.title = "Template"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
