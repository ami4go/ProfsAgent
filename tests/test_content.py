"""Phase 2 deliverable B (content/): deck checks, rendering and generate_material, offline with a fake LLM and the
committed clean CS-601 design (writes only to a temporary folder)."""
import json
import sys
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import pytest
from pptx import Presentation

from profsagent.content import generate as G
from profsagent.content.checks import check_deck
from profsagent.content.models import Figure, LectureDeck, Quiz, Slide
from profsagent.content.render import fit
from profsagent.delivery.brief import build_briefs

FIXTURE = ROOT / "tests/fixtures/cs601_design_state.json"


@pytest.fixture(scope="module")
def st():
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def brief(st):
    return build_briefs(st)[0]                                    # W01.L1, 2 h


@pytest.fixture
def lecture(brief):
    return G.lecture_input(brief)


def good_deck(lecture: dict, n: int = 18) -> LectureDeck:
    los = [lo["id"] for lo in lecture["learning_outcomes"]]
    slides = [Slide(kind="concept", title=f"Idea {i}", bullets=["one point", "another point"], notes="say this",
                    los=[los[i % len(los)]], sources=["C1"], minutes=lecture["minutes"] / n) for i in range(n - 4)]
    slides += [Slide(kind="worked_problem", title="Worked", bullets=["step 1", "step 2"], los=los, notes="n", minutes=lecture["minutes"] / n,
                     figure=Figure(kind="flow", steps=["Spec", "Prompt", "Test"])),
               Slide(kind="check", title="Check 1", quiz=Quiz(question="Which?", options=["a", "b", "c"], answer=1, explanation="b"),
                     los=los[:1], notes="n", minutes=lecture["minutes"] / n),
               Slide(kind="check", title="Check 2", quiz=Quiz(question="Why?", options=["x", "y", "z"], answer=0, explanation="x"),
                     los=los[:1], notes="n", minutes=lecture["minutes"] / n),
               Slide(kind="summary", title="Summary", bullets=["recap"], los=los, notes="n", minutes=lecture["minutes"] / n)]
    return LectureDeck(lecture_id=lecture["lecture_id"], title="T", slides=slides)


def codes(found, sev=None):
    return {x["code"] for x in found if sev is None or x["severity"] == sev}


CFG = G.cfg()


# ------------------------------------------------------------------ checks
def test_a_good_deck_has_no_errors(lecture):
    assert codes(check_deck(good_deck(lecture), lecture, {"C1"}, CFG), "error") == set()


def test_every_learning_outcome_must_be_served(lecture):
    deck = good_deck(lecture)
    lo = lecture["learning_outcomes"][0]["id"]
    for s in deck.slides:
        s.los = [x for x in s.los if x != lo]
    assert any(x["code"] == "D-LO" and lo in x["nodes"] for x in check_deck(deck, lecture, {"C1"}, CFG))


def test_unknown_ids_holes_and_bad_figures_are_errors(lecture):
    deck = good_deck(lecture)
    deck.slides[0].sources, deck.slides[1].los = ["L9"], ["CO9.LO9"]
    deck.slides[2].bullets = ["for {n} of {m} items"]
    deck.slides[3].figure = Figure(kind="flow", steps=["only one"])
    deck.slides[4].figure = Figure(kind="table", rows=[["a", "b"], ["c"]])
    assert {"D-REF", "D-TEMPLATE", "D-FIGURE"} <= codes(check_deck(deck, lecture, {"C1"}, CFG), "error")


def test_slide_count_follows_the_lecture_hours(lecture):
    assert "D-COUNT" in codes(check_deck(good_deck(lecture, n=6), lecture, {"C1"}, CFG), "error")   # 2 h: 16-28


def test_check_slides_need_a_valid_question(lecture):
    deck = good_deck(lecture)
    deck.slides[-3].quiz = None
    deck.slides[-2].quiz.answer = 7
    found = check_deck(deck, lecture, {"C1"}, CFG)
    assert sum(x["code"] == "D-QUIZ" and x["severity"] == "error" for x in found) == 2


def test_fewer_checks_than_hours_is_a_warning(lecture):
    deck = good_deck(lecture)
    deck.slides[-3].kind, deck.slides[-3].bullets = "concept", ["x"]
    assert "D-CHECK" in codes(check_deck(deck, lecture, {"C1"}, CFG), "warning")


# ------------------------------------------------------------------ rendering
def test_fit_shrinks_long_text_and_reports_what_cannot_fit():
    from pptx.util import Inches
    assert fit(["short"], Inches(6), Inches(1), 24, 14) == (24, True)
    size, ok = fit(["word " * 80], Inches(6), Inches(1.5), 24, 14)
    assert size < 24 or not ok
    assert fit(["word " * 400], Inches(3), Inches(0.5), 24, 14) == (14, False)


# ------------------------------------------------------------------ generate_material
class FakeLLM:
    def __init__(self, decks):
        self.decks, self.calls = decks, []

    def structured(self, *, system, user, model_cls, role, temperature, prompt_id, prompt_version, run_id=None):
        self.calls.append(user)
        return self.decks.pop(0), SimpleNamespace(endpoint="fake:model", prompt_tokens=10, output_tokens=20, latency_s=0.1)


@pytest.fixture
def run(st, tmp_path, monkeypatch):
    monkeypatch.setattr(G, "load_state", lambda run_id: st)
    monkeypatch.setattr(G, "out_dir", lambda run_id: tmp_path)
    return tmp_path


def test_generate_writes_slides_notes_quiz_and_manifest(run, lecture):
    llm = FakeLLM([good_deck(lecture)])
    m = G.generate_material("cs601", lecture["lecture_id"], llm=llm, log=lambda _: None)
    d = run / "material" / lecture["lecture_id"]
    assert len(llm.calls) == 1 and m["errors"] == 0 and m["cited"] == ["C1"]
    prs = Presentation(str(d / "slides.pptx"))
    assert len(prs.slides) == m["slides"] == 18 + 3                       # title + objectives + 18 + references
    assert "Answer: B. b" in prs.slides[2 + 15].notes_slide.notes_text_frame.text
    assert json.loads((d / "quiz.json").read_text(encoding="utf-8"))[0]["answer"] == 1
    notes = (d / "notes.md").read_text(encoding="utf-8")
    assert "## References" in notes and "Draft generated by ProfsAgent" in notes
    assert '"C1"' in llm.calls[0] and "Write the content slides" in llm.calls[0]


def test_a_deck_with_errors_is_repaired_once_and_kept_only_if_better(run, lecture):
    bad = good_deck(lecture, n=6)                                          # D-COUNT
    llm = FakeLLM([bad, good_deck(lecture)])
    m = G.generate_material("cs601", lecture["lecture_id"], llm=llm, log=lambda _: None)
    assert len(llm.calls) == 2 and "D-COUNT" in llm.calls[1] and m["errors"] == 0
    llm = FakeLLM([bad, good_deck(lecture, n=5)])                           # the repair is no better: keep the first
    m = G.generate_material("cs601", lecture["lecture_id"], llm=llm, log=lambda _: None)
    assert m["content_slides"] == 6 and m["errors"] == 1


def test_sources_are_course_resources_then_approved_lecture_resources(st, brief, run):
    lec = {"resources": {"items": {"reading": [{"title": "Approved", "url": "https://a.org", "kind": "reading", "status_decision": "approve"},
                                               {"title": "Pending", "url": "https://p.org", "kind": "reading"}]}}}
    (run / "lectures").mkdir()
    (run / "lectures" / f"{brief['lecture_id']}.json").write_text(json.dumps(lec), encoding="utf-8")
    src = G.gather_sources(st, "cs601", brief, include_pending=False, max_external=8)
    assert [s["title"] for k, s in src.items() if k.startswith("L")] == ["Approved"]
    assert all(s["approved"] for s in src.values()) and any(k.startswith("C") for k in src)
    assert len([k for k in G.gather_sources(st, "cs601", brief, True, 8) if k.startswith("L")]) == 2
