"""Offline tests for deterministic parts: scheduler, CO–PO computation, validators, store freezing."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import pytest

from profsagent.design import deterministic as det
from profsagent.design.store import DesignGraph, FrozenError
from profsagent.validate import rules


def _state():
    cos = [{"id": "CO1", "statement": "Implement a class hierarchy for a given specification, such that 80% of tests pass.",
            "verb": "implement", "behaviour": "a class hierarchy", "condition": "for a given specification",
            "degree": "such that 80% of tests pass", "bloom_level": 3, "knowledge_dimension": "procedural",
            "covers_topic_groups": ["TG1"], "evidence_types": ["programming_assignment"]},
           {"id": "CO2", "statement": "Understand and design patterns effectively.", "verb": "understand",
            "behaviour": "patterns", "condition": "", "degree": "effectively", "bloom_level": 6,
            "knowledge_dimension": "conceptual", "covers_topic_groups": [], "evidence_types": []}]
    los = [{"id": "CO1.LO1", "parent_co": "CO1", "statement": "Write a class for a given interface, with all 5 tests passing.",
            "verb": "write", "behaviour": "a class", "condition": "for a given interface", "degree": "with all 5 tests passing",
            "bloom_level": 3, "knowledge_dimension": "procedural", "requires": [], "hands_on": True, "est_lecture_h": 3}]
    modules = [{"id": "M1", "title": "OOP", "order": 1, "primary_cos": ["CO1"], "topics": [
        {"id": "M1.T1", "title": "classes", "est_hours": 2, "serves_los": ["CO1.LO1"], "requires": []},
        {"id": "M1.T2", "title": "orphan", "est_hours": 3, "serves_los": [], "requires": ["M1.T1"]}]}]
    sched = det.schedule([{**t, "module_order": 1} for t in modules[0]["topics"]], 13, 3, 7)
    assessment = {"components": [
        {"type": "quiz", "weight_pct": 10, "best_k": 2, "n": 2, "activity_types": [], "instances": [
            {"id": "A:Quiz1", "release_week": 1, "due_week": 1, "covers_topics": ["M1.T2"], "assesses": [{"lo_id": "CO1.LO1", "marks_pct_of_instance": 100, "bloom_level": 2}]},
            {"id": "A:Quiz2", "release_week": 3, "due_week": 3, "covers_topics": [], "assesses": [{"lo_id": "CO1.LO1", "marks_pct_of_instance": 100, "bloom_level": 3}]}]},
        {"type": "project", "weight_pct": 80, "activity_types": ["teamwork"], "instances": [
            {"id": "A:Proj", "release_week": 5, "due_week": 13, "covers_topics": [], "assesses": [{"lo_id": "CO1.LO1", "marks_pct_of_instance": 100, "bloom_level": 3}]}]}]}
    return {"cos": cos, "los": los, "modules": modules, "schedule": sched, "assessment": assessment, "weeks": 13,
            "bloom_band": {"floor": 2, "ceiling": 5}, "topic_groups": [{"id": "TG1"}, {"id": "TG2"}],
            "hours_budget": {"lecture_hours_total": 39}, "midsem_after_week": 7, "constraints": [], "ltp": {"L": 3, "T": 1, "P": 0}}


def codes(vs):
    return {x["code"] for x in vs if x["severity"] == "error"}


def test_schedule_splits_and_orders():
    s = det.schedule([{"id": "B", "est_hours": 2, "requires": ["A"]}, {"id": "A", "est_hours": 2.5, "requires": []}], 13, 3, 7)
    assert s["order"] == ["A", "B"] and s["feasible"]
    assert s["topic_weeks"]["B"] == [1, 2]


def test_schedule_infeasible_explains():
    s = det.schedule([{"id": f"T{i}", "est_hours": 3} for i in range(15)], 13, 3, 7)
    assert not s["feasible"] and "45.0 lecture hours" in s["explanation"]


def test_cycle_broken():
    s = det.schedule([{"id": "A", "est_hours": 1, "requires": ["B"], "suggested_order": 1},
                      {"id": "B", "est_hours": 1, "requires": ["A"], "suggested_order": 2}], 13, 3, 7)
    assert s["dropped_edges"] and len(s["order"]) == 2


def test_validators_catch_seeded_defects():
    st = _state()
    c = codes(rules.run(st))
    assert {"V12", "V11", "V7", "VCO-COVER", "V4"} <= c          # banned verb, vague degree, above band, TG2 uncovered, 2 COs
    assert "T5" in c                                              # orphan topic
    assert "V3" in c                                              # best 2 of 2
    assert "V19" in c                                             # Quiz1 (week 1) covers M1.T2, which runs into week 2
    assert "V1" in c                                              # 90 != 100
    assert "T2" in c                                              # CO2 has no LOs


def test_copo_computed_strength():
    st = _state()
    lo_parent = {"CO1.LO1": "CO1"}
    r = det.copo_strengths(["CO1", "CO2"], lo_parent, st["assessment"]["components"], {"teamwork": ["PO5"]},
                           [{"co_id": "CO1", "po_uid": "PO5", "strength_provisional": 3}, {"co_id": "CO2", "po_uid": "PO4", "strength_provisional": 3}],
                           {"three": 0.6, "two": 0.3})
    d = {(x["co_id"], x["po_uid"]): x for x in r}
    assert d[("CO1", "PO5")]["strength_computed"] == 3          # 80/(80+10) of CO1 marks from teamwork
    assert d[("CO2", "PO4")]["strength_computed"] == 0 and d[("CO2", "PO4")]["claimed"]
    vs = rules.check_copo_computed({"copo_computed": r})
    assert any(x["code"] == "V23" and x["severity"] == "error" for x in vs)


def test_frozen_nodes_cannot_change(tmp_path):
    g = DesignGraph("P1", tmp_path / "g.json")
    g.put("CO", "CO1", {"statement": "x"}, created_by="G4")
    g.freeze_labels(["CO"])
    with pytest.raises(FrozenError):
        g.put("CO", "CO1", {"statement": "y"}, created_by="R1")
    with pytest.raises(FrozenError):
        g.update("CO1", "statement", "z", by="R1")


def test_new_rules_catch_review_defects():
    st = _state()
    st["los"].append({"id": "CO1.LO2", "parent_co": "CO1", "statement": st["cos"][0]["statement"], "verb": "implement",
                      "behaviour": "x", "condition": "for a given specification", "degree": "such that 80% of tests pass",
                      "bloom_level": 3, "knowledge_dimension": "procedural"})
    st["los"].append({"id": "CO1.LO3", "parent_co": "CO1", "statement": "Write a parser for a given grammar, with 5 tests passing.",
                      "verb": "verify", "behaviour": "a parser", "condition": "for a given grammar", "degree": "with 5 tests passing",
                      "bloom_level": 5, "knowledge_dimension": "procedural"})
    st["assessment"]["components"][1]["instances"] = [
        {"id": "A:ProjM1", "release_week": 11, "due_week": 12, "covers_topics": [], "assesses": [{"lo_id": "CO1.LO1", "marks_pct_of_instance": 100, "bloom_level": 3}]},
        {"id": "A:ProjFinal", "release_week": 9, "due_week": 13, "covers_topics": [], "assesses": [{"lo_id": "CO1.LO1", "marks_pct_of_instance": 100, "bloom_level": 3}]}]
    c = codes(rules.run(st))
    assert "VLO-COPY" in c          # LO restates its CO
    assert "V10" in c               # statement verb 'Write' != declared 'verify'
    assert "VA-ORDER" in c          # milestone released after the final
    assert "VTG-TAUGHT" in c        # topic groups never taught (no covers_topic_groups)
    assert "V21" in c               # 5 h of topics in a 39 h budget (under-filled)


def test_fix_hints_attached():
    vs = rules.run(_state())
    assert all("fix_hint" in x for x in vs if x["code"] in ("V10", "V11", "V12", "T2"))


@pytest.mark.parametrize("level, code, years, expected", [
    (400, None, [3], 4),        # G1 wrote the course number instead of the level digit (crashed bloom_band before)
    (4, None, [3], 4),
    (None, "CSE401", [3], 4),
    (None, "CSE901", [3], 3),   # no Bloom band for level 9: fall back to the lowest target year
    (0, None, [], 3),
])
def test_course_level_normalised(level, code, years, expected):
    missing = {"status": "missing", "value": None}
    cco = {"fields": {"credits": missing, "semester_weeks": missing, "weekly_effort_hours": missing,
                      "ltp": {"status": "missing", "L": None, "T": None, "P": None}, "lab": {"required": "yes"},
                      "level_or_code": {"status": "ambiguous", "level": level, "code": code},
                      "target_students": {"years": years}}}
    d = det.apply_defaults(cco)
    assert d["level"] == expected
    assert det.bloom_band(d["level"])["floor"] >= 1


# ------------------------------------------------------------------ IIIT-D CSE course template rules
def test_template_co_count_and_high_bloom_rule():
    st = {"bloom_band": {"floor": 1, "ceiling": 6}, "co_count": det.co_count_for(4), "topic_groups": [],
          "cos": [{"id": f"CO{i}", "statement": f"Apply technique {i} to a given dataset, using Python, with 90% accuracy.",
                   "verb": "apply", "behaviour": f"technique {i}", "condition": "using Python", "degree": "90% accuracy",
                   "bloom_level": 3, "covers_topic_groups": []} for i in range(1, 7)]}
    codes = {x["code"] for x in rules.check_cos(st)}
    assert "V4" in codes and "VCO-BL" in codes          # 6 COs for 4 credits, none at BL4+
    assert det.co_count_for(2) == {"min": 2, "max": 3, "basis": "IIIT-D CSE course template (2-credit course)"}


def test_template_proctored_rule():
    comps = [{"type": "project", "weight_pct": 50, "instances": []}, {"type": "assignment", "weight_pct": 20, "instances": []},
             {"type": "midsem", "weight_pct": 30, "instances": []}]
    st = {"assessment": {"components": comps}, "los": [], "weeks": 13, "modules": [], "schedule": {"topic_weeks": {}}}
    assert "VA-PROCTOR" in {x["code"] for x in rules.check_assessment(st)}
    comps[1]["proctored"] = True                         # e.g. an invigilated in-lab test
    assert "VA-PROCTOR" not in {x["code"] for x in rules.check_assessment(st)}


def test_template_programmes_verbatim():
    from profsagent.config import pedagogy, programme
    sizes = {p: len(programme(p)["po"]) for p in pedagogy()["copo_programmes"]}
    assert sizes == {"BTECH-COMMON": 12, "BTECH-CSE": 5, "BTECH-CSAI": 6, "MTECH-CSE": 13}
    assert programme("BTECH-COMMON")["po"][4]["text"].startswith("Modern tool usage")


def test_template_xlsx_keeps_official_structure(tmp_path):
    import json

    import openpyxl
    from profsagent.render.course_template import write_xlsx
    src = ROOT / "runs" / "cs601-01" / "state.json"
    if not src.exists():
        pytest.skip("runs/cs601-01 not present (restore with: git restore runs/)")
    st = json.loads(src.read_text(encoding="utf-8"))
    ws = openpyxl.load_workbook(write_xlsx(st, tmp_path / "t.xlsx"))["Template"]
    col_a = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    for label in ("Code", "CourseOutcomes", "Approval Date", "EndMandatory", "Suggested Topics", "Suggested Plan of Labs",
                  "Suggested Assessment Plan", "Suggested Resources", "EndSuggested", "B.Tech. (common)", "B.Tech.(CSAI)"):
        assert label in col_a
    end = col_a.index("EndMandatory") + 1
    assert ws.cell(end, 2).value == "Do not remove or edit this row"
    assert [ws.cell(r, 1).value for r in range(col_a.index("CourseOutcomes") + 3, end - 2)] == [c["id"] for c in st["cos"]]


def test_fit_assessment_fixes_weights_and_timing():
    comps = [
        {"type": "assignment", "weight_pct": 35, "proctored": False,
         "instances": [{"id": "A1", "release_week": 2, "due_week": 4, "covers_topics": ["T3"],
                        "assesses": [{"lo_id": "CO1.LO1", "marks_pct_of_instance": 100, "bloom_level": 3}]}]},
        {"type": "project", "weight_pct": 45, "instances": [{"id": "P1", "release_week": 5, "due_week": 12, "covers_topics": [],
                                                              "assesses": [{"lo_id": "CO2.LO1", "marks_pct_of_instance": 100, "bloom_level": 6}]}]},
        {"type": "midsem", "weight_pct": 22, "instances": [{"id": "M", "release_week": 6, "due_week": 6, "covers_topics": ["T3", "T9"],
                                                             "assesses": [{"lo_id": "CO1.LO1", "marks_pct_of_instance": 50, "bloom_level": 3},
                                                                          {"lo_id": "CO3.LO1", "marks_pct_of_instance": 50, "bloom_level": 3}]}]},
    ]
    topic_last = {"T3": 3, "T9": 9}
    lo_last = {"CO1.LO1": 3, "CO3.LO1": 9}
    out, notes = det.fit_assessment(comps, topic_last, lo_last, weeks=13, midsem_after=7,
                                    proctored={"min_pct": 40, "types": ["midsem", "endsem", "quiz"]})
    assert sum(c["weight_pct"] for c in out) == 100                        # 102 → exactly 100
    assert out[2]["weight_pct"] >= 40                                       # proctored share raised to the minimum
    a1 = out[0]["instances"][0]
    assert (a1["release_week"], a1["due_week"]) == (3, 5)                  # released after T3 is taught, duration kept
    m = out[2]["instances"][0]
    assert m["release_week"] == 7 and m["covers_topics"] == ["T3"]         # exam stays on the calendar, drops untaught T9
    assert [a["lo_id"] for a in m["assesses"]] == ["CO1.LO1"] and m["assesses"][0]["marks_pct_of_instance"] == 100
    assert notes


def test_design_graph_loads_into_neo4j(tmp_path):
    """Round trip: a tiny design run loads into Neo4j, links to a real Layer A course, then is removed again."""
    import json

    from profsagent.design.neo4j_sync import delete_run, driver, load_run
    try:
        drv = driver()
        drv.verify_connectivity()
    except Exception:
        pytest.skip("Neo4j not reachable")
    g = DesignGraph("TST-999", tmp_path / "design_graph.json")
    g.put("NewCourse", "COURSE", {"title": "Loader test"}, created_by="G1")
    g.put("CO", "CO1", {"statement": "Design X."}, created_by="G4")
    g.put("LO", "CO1.LO1", {"statement": "Build Y."}, created_by="G6")
    g.link("CO1", "DECOMPOSES_INTO", "CO1.LO1")
    g.link("COURSE", "REQUIRES", "CSE201")
    g.link("CO1", "MAPS_TO", "BTECH-COMMON/PO3")
    run = tmp_path / "pytest-loader-run"
    run.mkdir()
    g.path = run / "design_graph.json"
    g.save()
    (run / "state.json").write_text(json.dumps({"cco_compact": {"course_name": "Loader test"}}), encoding="utf-8")
    try:
        stats = load_run(run, drv)
        assert stats["nodes"] == 3 and stats["internal_links"] == 1 and stats["links_to_programme"] == 1
        with drv.session() as s:
            assert s.run("MATCH (:D_NewCourse {run_id: $r})-[:REQUIRES]->(c:Course {code: 'CSE201'}) RETURN count(*) AS n",
                         r=run.name).single()["n"] == 1
    finally:
        assert delete_run(run.name, drv) == 3
        drv.close()


def test_g10_prompt_renders_with_and_without_replacement_round():
    from profsagent.llm.prompts import load_prompt
    base = dict(modules_topics=[], cco={"course_name": "x"}, corpus_resources=[], retrieved_external=[])
    assert "REPLACEMENT ROUND" not in load_prompt("G10").render(**base)[1]
    assert "REPLACEMENT ROUND" in load_prompt("G10").render(**base, rejected_picks=[{"title": "t", "problems": []}],
                                                            already_selected=[])[1]
